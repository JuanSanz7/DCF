import os
import sys
import re
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
from scipy.stats import spearmanr
from datetime import datetime
from matplotlib.gridspec import GridSpec # Import GridSpec
from typing import Dict, Optional

# ---------------------------------------------------------------------------
# Damodaran DCF template import (Google Sheets / Excel)
# ---------------------------------------------------------------------------
# Reads each input by matching the LABEL TEXT in column A (not a fixed row
# number), so it survives rows being inserted/deleted/reordered in a given
# company tab — as long as the label text itself isn't renamed. Matching is
# accent-insensitive, case-insensitive, substring-based, and takes the FIRST
# matching row top-to-bottom (relevant for "Precio de mercado actual", which
# appears both as the input and again as a later reference/output).
LABEL_MAP = {
    "risk_free_rate": "tasa libre de riesgo (rf)",
    "equity_risk_premium": "prima de riesgo (erp)",
    "tax_rate": "tasa impositiva (usada",
    "debt": "deuda total - debt (valor contable)",
    "WACC": "cost of capital (wacc)",
    "shares": "acciones en circulacion, diluidas (m)",
    "price": "precio de mercado actual",
    "ebit": "ebit (con i+d capitalizado) - base del modelo",
    "cash": "caja y equivalentes (+ inversiones corto plazo)",
    "growth_rate_5y": "growth 5y - crecimiento nopat, anos 1-5",
    "growth_rate_5_10y": "growth 5-10y - crecimiento nopat, anos 6-10",
    "reinvestment_rate_5y": "reinvestment 5y - tasa de reinversion, anos 1-5",
    "reinvestment_rate_5_10y": "reinvestment 5-10y - tasa de reinversion, anos 6-10",
}
# These fields are stored as fractions in the template (e.g. 0.135 = 13.5%)
# and must be converted to the app's percentage-point convention (13.5).
_PERCENT_FIELDS = {
    "risk_free_rate", "equity_risk_premium", "tax_rate", "WACC",
    "growth_rate_5y", "growth_rate_5_10y", "reinvestment_rate_5y", "reinvestment_rate_5_10y",
}
# Fields the app (Streamlit and Tkinter GUI) actually consumes.
APP_FIELDS = [
    "price", "shares", "cash", "ebit", "debt", "tax_rate",
    "growth_rate_5y", "growth_rate_5_10y", "risk_free_rate", "equity_risk_premium",
    "WACC", "reinvestment_rate_5y", "reinvestment_rate_5_10y",
]


def extract_sheet_id(url_or_id: str) -> str:
    """Pulls the spreadsheet ID out of a full Google Sheets URL, or returns
    the input unchanged if it's already a bare ID."""
    m = re.search(r"/d/([a-zA-Z0-9_-]+)", url_or_id)
    return m.group(1) if m else url_or_id.strip()


def _normalize(s) -> str:
    """Lowercases and strips accents, so 'Año' / 'Ano' / 'AÑO' all match."""
    import unicodedata
    s = unicodedata.normalize("NFKD", str(s))
    s = "".join(c for c in s if not unicodedata.combining(c))
    return s.lower().strip()


def _clean_numeric(raw) -> Optional[float]:
    """Parses a cell value that may be formatted as '13.50%', '1,234.5', or
    a plain number. Returns None if it can't be parsed."""
    if raw is None:
        return None
    if isinstance(raw, (int, float)) and not (isinstance(raw, float) and pd.isna(raw)):
        return float(raw)
    s = str(raw).strip()
    if not s or s.lower() == "nan":
        return None
    was_percent = s.endswith("%")
    s = s.replace("%", "").replace(",", "").replace("€", "").replace("$", "").strip()
    try:
        val = float(s)
    except ValueError:
        return None
    return val / 100.0 if was_percent else val


def _find_row_by_label(column_a_values, label_substring: str) -> Optional[int]:
    """Returns the 0-indexed position of the first cell in column_a_values
    whose normalized text contains label_substring (already normalized)."""
    for i, cell in enumerate(column_a_values):
        if cell is None:
            continue
        if label_substring in _normalize(cell):
            return i
    return None


def _grid_to_app_values(column_a_values, get_col_b) -> Dict:
    """Shared logic: given the full column-A text (for label search) and a
    function get_col_b(row_index_0based) -> raw column-B value, builds the
    dict of app-ready values by matching each field's label text."""
    data = {}
    missing_fields = []
    for key in APP_FIELDS:
        label = LABEL_MAP.get(key)
        row_idx = _find_row_by_label(column_a_values, label) if label else None
        raw = get_col_b(row_idx) if row_idx is not None else None
        val = _clean_numeric(raw)
        if val is None:
            missing_fields.append(key)
            continue
        data[key] = val * 100.0 if key in _PERCENT_FIELDS else val
    data["_read_ok"] = True
    data["_missing_fields"] = missing_fields
    return data


def fetch_from_google_sheet(sheet_url: str, tab_name: str) -> Dict:
    """Reads a company tab from a PUBLIC Google Sheet built from the
    Damodaran DCF template (no credentials required — the tab must be
    shared as 'Anyone with the link can view'). Matches inputs by their
    label text in column A, so it's robust to rows being inserted, deleted,
    or reordered relative to the original template.
    """
    sheet_id = extract_sheet_id(sheet_url)
    tab_encoded = tab_name.replace(" ", "%20")
    csv_url = f"https://docs.google.com/spreadsheets/d/{sheet_id}/gviz/tq?tqx=out:csv&sheet={tab_encoded}"
    df = pd.read_csv(csv_url, header=None)
    column_a = df.iloc[:, 0].tolist() if df.shape[1] > 0 else []

    def get_col_b(row_idx):
        try:
            return df.iloc[row_idx, 1] if df.shape[1] > 1 else None
        except (IndexError, KeyError):
            return None

    return _grid_to_app_values(column_a, get_col_b)


def fetch_from_excel_template(file_like, tab_name: str) -> Dict:
    """Reads a company tab from a .xlsx workbook built from the Damodaran
    DCF template (use this for private workbooks — no sharing settings
    needed). `file_like` is a file path or a file-like object, e.g.
    Streamlit's UploadedFile from st.file_uploader. Matches inputs by their
    label text in column A, same as fetch_from_google_sheet.
    """
    import openpyxl
    wb = openpyxl.load_workbook(file_like, data_only=True)
    if tab_name not in wb.sheetnames:
        raise ValueError(f"Tab '{tab_name}' not found. Available tabs: {', '.join(wb.sheetnames)}")
    ws = wb[tab_name]
    column_a = [ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)]

    def get_col_b(row_idx):
        return ws.cell(row=row_idx + 1, column=2).value

    return _grid_to_app_values(column_a, get_col_b)


# ---------------------------------------------------------------------------
# Monte Carlo simulation engine
# ---------------------------------------------------------------------------
def run_monte_carlo_simulation(params):
    # Unpack parameters
    company_name = params['company_name']
    currency = params.get('currency', 'USD')
    current_price = params['current_price']
    operating_income_base = params['operating_income_base']  # This is EBIT
    tax_rate = params.get('tax_rate', 0.21)  # Default to 21% if not provided
    shares_outstanding = params['shares_outstanding']
    cash = params['cash']
    debt = params['debt']
    growth_rate_5y = params['growth_rate_5y']
    growth_rate_5_10y = params['growth_rate_5_10y']
    risk_free_rate = params['risk_free_rate']
    equity_risk_premium = params['equity_risk_premium']
    WACC = params['WACC']
    reinvestment_rate_5y = params['reinvestment_rate_5y']
    reinvestment_rate_5_10y = params['reinvestment_rate_5_10y']
    std_growth_5y = params['std_growth_5y']
    std_growth_5_10y = params['std_growth_5_10y']
    std_risk_free = params['std_risk_free']
    std_equity_premium = params['std_equity_premium']
    std_WACC = params['std_WACC']
    std_reinv_5y = params['std_reinv_5y']
    std_reinv_5_10y = params['std_reinv_5_10y']
    n_simulations = params['n_simulations']
    current_date = datetime.now().strftime("%Y-%m-%d")
    
    # Calculate NOPAT from EBIT: NOPAT = EBIT * (1 - tax_rate)
    nopat_base = operating_income_base * (1 - tax_rate)
    
    # Calculate terminal value parameters using base values
    terminal_growth_base = risk_free_rate
    terminal_WACC_base = risk_free_rate + equity_risk_premium
    terminal_reinv_rate_base = risk_free_rate / (risk_free_rate + equity_risk_premium) if (risk_free_rate + equity_risk_premium) > 0 else 0

    np.random.seed(42)
    params_simulated = {
        'growth_5y': [],
        'growth_5_10y': [],
        'risk_free': [],
        'equity_premium': [],
        'WACC': [],
        'reinv_5y': [],
        'reinv_5_10y': [],
        'value_per_share': []
    }
    results = []
    fcf_projections = []
    for _ in range(n_simulations):
        growth_rate_5y_sim = np.random.normal(growth_rate_5y, std_growth_5y)
        growth_rate_5_10_sim = np.random.normal(growth_rate_5_10y, std_growth_5_10y)
        risk_free_rate_sim = np.random.normal(risk_free_rate, std_risk_free)
        equity_risk_premium_sim = np.random.normal(equity_risk_premium, std_equity_premium)
        WACC_sim = np.random.normal(WACC, std_WACC)
        reinv_rate_5y_sim = np.random.normal(reinvestment_rate_5y, std_reinv_5y)
        reinv_rate_5_10y_sim = np.random.normal(reinvestment_rate_5_10y, std_reinv_5_10y)
        params_simulated['growth_5y'].append(growth_rate_5y_sim)
        params_simulated['growth_5_10y'].append(growth_rate_5_10_sim)
        params_simulated['risk_free'].append(risk_free_rate_sim)
        params_simulated['equity_premium'].append(equity_risk_premium_sim)
        params_simulated['WACC'].append(WACC_sim)
        params_simulated['reinv_5y'].append(reinv_rate_5y_sim)
        params_simulated['reinv_5_10y'].append(reinv_rate_5_10y_sim)
        terminal_WACC = risk_free_rate_sim + equity_risk_premium_sim
        terminal_growth = risk_free_rate_sim
        reinvestment_rate_terminal = risk_free_rate_sim / (risk_free_rate_sim + equity_risk_premium_sim)
        # Use NOPAT instead of operating income for FCF calculations
        nopat = nopat_base
        FCFs = []
        nopats = []
        for year in range(1, 6):
            nopat *= (1 + growth_rate_5y_sim)
            nopats.append(nopat)
            FCF = nopat * (1 - reinv_rate_5y_sim)
            FCFs.append(FCF)
        for year in range(6, 11):
            nopat *= (1 + growth_rate_5_10_sim)
            nopats.append(nopat)
            FCF = nopat * (1 - reinv_rate_5_10y_sim)
            FCFs.append(FCF)
        nopat_terminal = nopat * (1 + terminal_growth)
        FCF_terminal = nopat_terminal * (1 - reinvestment_rate_terminal)
        terminal_value = FCF_terminal / (terminal_WACC - terminal_growth)
        fcf_projections.append(FCFs)
        discount_factors = [(1 + WACC_sim) ** i for i in range(1, 11)]
        PV_FCF = sum([f / d for f, d in zip(FCFs, discount_factors)])
        PV_terminal = terminal_value / ((1 + WACC_sim) ** 10)
        EV = PV_FCF + PV_terminal
        market_value = EV + cash - debt
        value_per_share = market_value / shares_outstanding
        results.append(value_per_share)
        params_simulated['value_per_share'].append(value_per_share)
    results = np.array(results)
    mean_value = np.mean(results)
    median_value = np.median(results)
    std_value = np.std(results)
    ci_lower = np.percentile(results, 2.5)
    ci_upper = np.percentile(results, 97.5)
    var_95 = np.percentile(results, 5)
    cvar_95 = np.mean(results[results < var_95])
    prob_overvalued = np.mean(results < current_price) * 100
    prob_undervalued = np.mean(results > current_price) * 100
    upside_potential = ((mean_value - current_price) / current_price) * 100
    df_params = pd.DataFrame(params_simulated)
    sensitivities = {}
    for param in df_params.columns[:-1]:
        correlation = spearmanr(df_params[param], df_params['value_per_share'])[0]
        std_norm = (df_params[param] - df_params[param].mean()) / df_params[param].std()
        coef = np.polyfit(std_norm, df_params['value_per_share'], 1)[0]
        sensitivities[param] = {'correlation': correlation, 'impact': coef}

    # Create the original combined 2x2 figure (fig_es)
    fig_es, axs = plt.subplots(2, 2, figsize=(15, 10))
    plt.style.use('seaborn-v0_8-darkgrid')

    # Plot 1: Intrinsic Value Distribution
    ax = axs[0, 0]
    ax.hist(results, bins=50, density=True, alpha=0.7, color='skyblue', edgecolor='black')
    ax.axvspan(mean_value - 3*std_value, mean_value - 2*std_value, color='red', alpha=0.1, label='±3σ')
    ax.axvspan(mean_value + 2*std_value, mean_value + 3*std_value, color='red', alpha=0.1)
    ax.axvspan(mean_value - 2*std_value, mean_value - std_value, color='orange', alpha=0.1, label='±2σ')
    ax.axvspan(mean_value + std_value, mean_value + 2*std_value, color='orange', alpha=0.1)
    ax.axvspan(mean_value - std_value, mean_value + std_value, color='green', alpha=0.1, label='±1σ')
    ax.axvline(mean_value, color='red', linestyle='--', label='Mean')
    ax.axvline(current_price, color='purple', linestyle='-', label='Current Price')
    ax.set_title(f'{company_name} - Intrinsic Value Distribution')
    ax.set_xlabel(f'Intrinsic Value per Share ({currency})')
    ax.set_ylabel('Density')
    ax.legend()

    # Plot 2: FCF Projection
    ax = axs[1, 0]
    fcf_mean = np.mean(fcf_projections, axis=0)
    fcf_std = np.std(fcf_projections, axis=0)
    years = range(1, 11)
    ax.fill_between(years, 
                    fcf_mean - 3*fcf_std,
                    fcf_mean + 3*fcf_std,
                    color='red',
                    alpha=0.1,
                    label='±3σ')
    ax.fill_between(years, 
                    fcf_mean - 2*fcf_std,
                    fcf_mean + 2*fcf_std,
                    color='orange',
                    alpha=0.1,
                    label='±2σ')
    ax.fill_between(years, 
                    fcf_mean - fcf_std,
                    fcf_mean + fcf_std,
                    color='green',
                    alpha=0.1,
                    label='±1σ')
    ax.plot(years, fcf_mean, marker='o', color='blue', label='Average FCF')
    ax.set_title(f'{company_name} - Free Cash Flow Projection')
    ax.set_xlabel('Year')
    ax.set_ylabel(f'FCF (Millions {currency})')
    ax.legend()
    ax.grid(True)

    # Plot 3: Sensitivity Analysis (Tornado plot) within fig_es
    ax = axs[1, 1]
    sensitivity_data = pd.DataFrame.from_dict(sensitivities, orient='index')
    sensitivity_data = sensitivity_data.sort_values('impact', ascending=True)
    ax.barh(range(len(sensitivity_data)), sensitivity_data['impact'], align='center')
    ax.set_yticks(range(len(sensitivity_data)))
    ax.set_yticklabels(sensitivity_data.index)
    ax.set_title(f'{company_name} - Sensitivity Analysis')
    ax.set_xlabel(f'Impact on Value per Share ({currency}/σ)')

    # Matplotlib-rendered Valuation Summary within fig_es
    axs[0, 1].axis('off')
    title = (
        f"VALUATION SUMMARY - {company_name}\n"
        f"Date: {current_date}\n"
        f"-------------------"
    )
    left_column = (
        f"Current Price: {current_price:.2f} {currency}\n"
        f"Mean Value: {mean_value:.2f} {currency}\n"
        f"Median Value: {median_value:.2f} {currency}\n"
        f"Upside Potential: {upside_potential:.1f}%\n\n"
        f"Probabilities:\n"
        f"Overvaluation: {prob_overvalued:.1f}%\n"
        f"Undervaluation: {prob_undervalued:.1f}%\n\n"
        f"Risk Metrics:\n"
        f"VaR 95%: {var_95:.2f} {currency}\n"
        f"CVaR 95%: {cvar_95:.2f} {currency}\n"
        f"Std. Deviation: {std_value:.2f} {currency}"
    )
    right_column = (
        f"Variable Parameters:\n"
        f"Growth 5y: {growth_rate_5y*100:.1f}% (±{std_growth_5y*100:.1f}%)\n"
        f"Growth 5-10y: {growth_rate_5_10y*100:.1f}% (±{std_growth_5_10y*100:.1f}%)\n"
        f"WACC: {WACC*100:.1f}% (±{std_WACC*100:.1f}%)\n"
        f"Risk Premium: {equity_risk_premium*100:.1f}% (±{std_equity_premium*100:.1f}%)\n"
        f"Risk Free Rate: {risk_free_rate*100:.1f}% (±{std_risk_free*100:.1f}%)\n"
        f"Reinvestment 5y: {reinvestment_rate_5y*100:.1f}% (±{std_reinv_5y*100:.1f}%)\n"
        f"Reinvestment 5-10y: {reinvestment_rate_5_10y*100:.1f}% (±{std_reinv_5_10y*100:.1f}%)\n"
        f"\nTerminal Value Params:\n"
        f"Term. Growth: {terminal_growth_base*100:.2f}%\n"
        f"Term. WACC: {terminal_WACC_base*100:.2f}%\n"
        f"Term. Reinv Rate: {terminal_reinv_rate_base*100:.2f}%"
    )
    left_lines = left_column.split('\n')
    right_lines = right_column.split('\n')
    max_lines = max(len(left_lines), len(right_lines))
    left_lines.extend([''] * (max_lines - len(left_lines)))
    right_lines.extend([''] * (max_lines - len(right_lines)))
    combined_lines = []
    for left, right in zip(left_lines, right_lines):
        combined_lines.append(f"{left:<35} {right:<35}")
    summary_text = title + '\n\n' + '\n'.join(combined_lines)
    axs[0, 1].text(0.5, 0.5, summary_text, fontsize=10, fontfamily='monospace', 
                  horizontalalignment='center', verticalalignment='center',
                  bbox=dict(facecolor='white', alpha=0.8))

    plt.tight_layout()

    # Create a separate figure for Intrinsic Value Distribution only
    fig_distribution_only, ax_dist_only = plt.subplots(figsize=(10, 6))
    plt.style.use('seaborn-v0_8-darkgrid')
    ax_dist_only.hist(results, bins=50, density=True, alpha=0.7, color='skyblue', edgecolor='black')
    ax_dist_only.axvspan(mean_value - 3*std_value, mean_value - 2*std_value, color='red', alpha=0.1, label='±3σ')
    ax_dist_only.axvspan(mean_value + 2*std_value, mean_value + 3*std_value, color='red', alpha=0.1)
    ax_dist_only.axvspan(mean_value - 2*std_value, mean_value - std_value, color='orange', alpha=0.1, label='±2σ')
    ax_dist_only.axvspan(mean_value + std_value, mean_value + 2*std_value, color='orange', alpha=0.1)
    ax_dist_only.axvspan(mean_value - std_value, mean_value + std_value, color='green', alpha=0.1, label='±1σ')
    ax_dist_only.axvline(mean_value, color='red', linestyle='--', label='Mean')
    ax_dist_only.axvline(current_price, color='purple', linestyle='-', label='Current Price')
    ax_dist_only.set_title(f'{company_name} - Intrinsic Value Distribution')
    ax_dist_only.set_xlabel(f'Intrinsic Value per Share ({currency})')
    ax_dist_only.set_ylabel('Density')
    ax_dist_only.legend()
    plt.tight_layout()

    # Create a separate figure for Sensitivity Analysis only
    fig_sensitivity, ax_sens_only = plt.subplots(figsize=(10, 6))
    sensitivity_data = pd.DataFrame.from_dict(sensitivities, orient='index')
    sensitivity_data = sensitivity_data.sort_values('impact', ascending=True)
    ax_sens_only.barh(range(len(sensitivity_data)), sensitivity_data['impact'], align='center')
    ax_sens_only.set_yticks(range(len(sensitivity_data)))
    ax_sens_only.set_yticklabels(sensitivity_data.index)
    ax_sens_only.set_title(f'{company_name} - Sensitivity Analysis')
    ax_sens_only.set_xlabel(f'Impact on Value per Share ({currency}/σ)')
    plt.tight_layout()

    # Create valuation summary dictionary
    valuation_summary = {
        'company_name': company_name,
        'date': current_date,
        'current_price': f"{current_price:.2f} {currency}",
        'mean_value': f"{mean_value:.2f} {currency}",
        'median_value': f"{median_value:.2f} {currency}",
        'upside_potential': f"{upside_potential:.1f}%",
        'prob_overvalued': f"{prob_overvalued:.1f}%",
        'prob_undervalued': f"{prob_undervalued:.1f}%",
        'VaR 95%': f"{var_95:.2f} {currency}",
        'CVaR 95%': f"{cvar_95:.2f} {currency}",
        'Std. Deviation': f"{std_value:.2f} {currency}",
        'Variable Parameters': {
            'Growth 5y': f"{growth_rate_5y*100:.1f}% (±{std_growth_5y*100:.1f}%)",
            'Growth 5-10y': f"{growth_rate_5_10y*100:.1f}% (±{std_growth_5_10y*100:.1f}%)",
            'WACC': f"{WACC*100:.1f}% (±{std_WACC*100:.1f}%)",
            'Risk Premium': f"{equity_risk_premium*100:.1f}% (±{std_equity_premium*100:.1f}%)",
            'Risk Free Rate': f"{risk_free_rate*100:.1f}% (±{std_risk_free*100:.1f}%)",
            'Reinvestment 5y': f"{reinvestment_rate_5y*100:.1f}% (±{std_reinv_5y*100:.1f}%)",
            'Reinvestment 5-10y': f"{reinvestment_rate_5_10y*100:.1f}% (±{std_reinv_5_10y*100:.1f}%)"
        },
        'Terminal Value Params': {
            'Term. Growth': f"{terminal_growth_base*100:.2f}%",
            'Term. WACC': f"{terminal_WACC_base*100:.2f}%",
            'Term. Reinv Rate': f"{terminal_reinv_rate_base*100:.2f}%"
        }
    }

    return fig_es, fig_distribution_only, fig_sensitivity, valuation_summary
